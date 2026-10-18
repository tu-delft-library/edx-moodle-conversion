import logging
import posixpath
import re
import shutil
import tempfile
import zipfile
from dataclasses import dataclass, field
from html import escape
from pathlib import Path, PurePosixPath
from urllib.parse import urlsplit

from bs4 import BeautifulSoup, NavigableString
from docx import Document

from ocw.assets import STATIC_REF_RE, WIDGET_MEDIA_LINKS, asset_name, is_skipped_asset
from ocw.converter.external_report import (
    ExternalContentReport,
    ExternalRef,
    _is_blacklisted,
    _OCW_HOST,
    _in_domain,
    _local_kind,
    _refs_in_html,
)
from ocw.converter.html_to_docx import HtmlToDocx
from ocw.parser.base import BaseParser
from ocw.utils import _UNSAFE_FILENAME_RE, PLATFORM_HOST_RE

log = logging.getLogger("ocw.zip_builder")
_ROOT_INDEX = PurePosixPath("index.html")
_COURSE_DOCX = "course.docx"
_EXTERNAL_XLSX = "external_content.xlsx"
_CONTENT_HEADING_OFFSET = 4
_PAGE = """<!doctype html>
<html lang="en">
<head><meta charset="utf-8"><title>{title}</title></head>
<body>
{nav}{body}
</body>
</html>
"""
_EMPTY_BODY = "<p>No content was found on the source page.</p>"
_COURSE_OVERVIEW = "course overview"
_NOT_REFERENCED = "not referenced on any page"
_TEACHERS_SEGMENT = "teachers"
_FOOTNOTE_FRAGMENT_RE = re.compile(r"^#(?:fn|fnref)\d+$")


def _norm(url: str) -> str:
    """Host and path of `url`, scheme, query, fragment and trailing slash ignored."""
    parts = urlsplit(url.strip())
    return f"{(parts.hostname or '').lower()}{parts.path.rstrip('/')}"


def _is_dead_host(host: str) -> bool:
    return _in_domain(host, _OCW_HOST) or bool(PLATFORM_HOST_RE.search(host))


def _is_dead_footnote(url: str) -> bool:
    parts = urlsplit(url.strip())
    return _is_dead_host(parts.hostname or "") and bool(_FOOTNOTE_FRAGMENT_RE.match(f"#{parts.fragment}"))


def _key(url: str) -> str:
    return url.split("#")[0].rstrip("/")


def _safe_name(text: str, index: int) -> str:
    cleaned = _UNSAFE_FILENAME_RE.sub("_", text).strip("._") or "untitled"
    return f"{index:02d}_{cleaned[:60]}"


def _rel(target: PurePosixPath, page: PurePosixPath) -> str:
    return posixpath.relpath(target.as_posix(), start=page.parent.as_posix())


def _folder(page: PurePosixPath) -> str:
    return f"{page.parent.as_posix()}/"


@dataclass
class PlannedPage:
    title: str
    path: PurePosixPath
    components: list[dict] = field(default_factory=list)
    summary_html: str = ""


@dataclass
class PlannedChapter:
    index: PlannedPage
    sequentials: list[tuple[str, list[PlannedPage]]] = field(default_factory=list)


@dataclass
class Plan:
    chapters: list[PlannedChapter] = field(default_factory=list)
    readings: list[PlannedPage] = field(default_factory=list)
    shared: list[PlannedPage] = field(default_factory=list)

    def content_pages(self) -> list[PlannedPage]:
        pages = []
        for chapter in self.chapters:
            for _, verticals in chapter.sequentials:
                pages.extend(verticals)
        return pages + self.readings + self.shared


class ZipBuilder:
    """Write a parsed course as a browsable folder structure inside a zip archive."""

    def __init__(self, course: BaseParser) -> None:
        self.course = course
        self._by_url: dict[str, PlannedPage] = {}
        self._pages: dict[str, PlannedPage] = {}
        self._teachers: PlannedPage | None = None
        self._home = PlannedPage(_COURSE_OVERVIEW, _ROOT_INDEX)
        self._skipped_assets: set[str] = set()
        self._usage: dict[str, list[str]] = {}

    def build(self, out: Path) -> None:
        tmp = Path(tempfile.mkdtemp())
        try:
            slug = _safe_name(self.course.course_name, 0)[3:]
            root = tmp / slug
            plan = self._plan()
            self._copy_assets(root)
            self._write_pages(root, plan)
            self._write_index(root, plan)
            report = ExternalContentReport(self._external_refs(plan) + self._local_refs())
            self._write_docx(root, plan, report)
            report.write_xlsx(root / _EXTERNAL_XLSX)
            self._zip(tmp, out)
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def _write_docx(self, root: Path, plan: Plan, report: ExternalContentReport) -> None:
        doc = Document()
        report.add_to(doc)
        self._add_contents(doc, plan)
        with tempfile.TemporaryDirectory() as scratch:
            HtmlToDocx(doc, root / "assets", Path(scratch)).add(
                self._flatten(plan), heading_offset=_CONTENT_HEADING_OFFSET
            )
            doc.save(root / _COURSE_DOCX)

    @staticmethod
    def _add_contents(doc, plan: Plan) -> None:
        doc.add_heading("Contents", level=1)
        for chapter in plan.chapters:
            doc.add_paragraph(chapter.index.title, style="List Bullet")
            for title, _ in chapter.sequentials:
                doc.add_paragraph(title, style="List Bullet 2")

    def _flatten(self, plan: Plan) -> str:
        """The whole course as one linear HTML document, structure headings marked `data-structural`."""
        parts = [f"<h1 data-structural>{escape(self.course.course_name)}</h1>"]
        parts += self._banner()
        if self.course.overview_summary_html:
            parts.append(self._localise(self.course.overview_summary_html, _ROOT_INDEX))
        for chapter in plan.chapters:
            parts.append(f"<h2 data-structural data-page-break>{escape(chapter.index.title)}</h2>")
            if chapter.index.summary_html:
                parts.append(self._localise(chapter.index.summary_html, chapter.index.path, _ROOT_INDEX))
            for title, verticals in chapter.sequentials:
                parts.append(f"<h3 data-structural>{escape(title)}</h3>")
                for vertical in verticals:
                    parts.append(f"<h4 data-structural>{escape(vertical.title)}</h4>")
                    parts.append(self._body_docx(vertical.components, vertical.path) or _EMPTY_BODY)
        return "\n".join(parts)

    def _body_docx(self, components: list[dict], page: PurePosixPath, stack: tuple = ()) -> str:
        return "\n".join(self._render_docx(comp, page, stack) for comp in components)

    def _render_docx(self, comp: dict, page: PurePosixPath, stack: tuple = ()) -> str:
        """Like `_render`, but a linked page is written in full where it is linked and a video is only a link."""
        match comp["type"]:
            case "video":
                return self._video_link(comp)
            case "reading_link" | "dedup_link":
                url = comp["reading_url"] if comp["type"] == "reading_link" else comp["dedup_url"]
                target = self._by_url[_key(url)]
                if target.path in stack:
                    return f"<p>{escape(target.title)} (see above)</p>"
                body = self._body_docx(target.components, target.path, (*stack, target.path))
                return f"<h5 data-structural>{escape(target.title)}</h5>\n{body}"
            case _:
                return self._render(comp, page, _ROOT_INDEX)

    def _plan(self) -> Plan:
        plan = Plan()
        for i, chapter in enumerate(self.course.chapters, 1):
            chapter_dir = PurePosixPath(_safe_name(chapter["display_name"], i))
            index = PlannedPage(
                chapter["display_name"],
                chapter_dir / "index.html",
                summary_html=chapter.get("summary_html", ""),
            )
            planned = PlannedChapter(index)
            self._register(chapter.get("source_url"), index)
            if chapter.get("source_url") and _norm(chapter["source_url"]).endswith(f"/{_TEACHERS_SEGMENT}"):
                self._teachers = index
            for j, seq in enumerate(chapter["sequentials"], 1):
                seq_dir = chapter_dir / _safe_name(seq["display_name"], j)
                verticals = [
                    PlannedPage(
                        vert["display_name"],
                        seq_dir / f"{_safe_name(vert['display_name'], k)}.html",
                        vert["components"],
                    )
                    for k, vert in enumerate(seq["verticals"], 1)
                ]
                for vert, page in zip(seq["verticals"], verticals):
                    self._register(vert.get("source_url"), page)
                planned.sequentials.append((seq["display_name"], verticals))
            plan.chapters.append(planned)
        self._plan_linked(plan.readings, self.course.reading_pages, "readings", "reading_url")
        self._plan_linked(plan.shared, self.course.dedup_pages, "shared", "dedup_url")
        return plan

    def _external_refs(self, plan: Plan) -> list[ExternalRef]:
        refs: list[ExternalRef] = []

        def _add(folder: str, components: list[dict]) -> None:
            refs.extend(self._component_refs(folder, components))

        if self.course.overview_summary_html:
            _add(_COURSE_OVERVIEW, [{"type": "html", "content": self.course.overview_summary_html}])
        for chapter in plan.chapters:
            if chapter.index.summary_html:
                _add(_folder(chapter.index.path), [{"type": "html", "content": chapter.index.summary_html}])
            for _, verticals in chapter.sequentials:
                for vertical in verticals:
                    _add(_folder(vertical.path), vertical.components)
        for page in plan.readings + plan.shared:
            _add(_folder(page.path), page.components)
        return refs

    @staticmethod
    def _component_refs(folder: str, components: list[dict]) -> list[ExternalRef]:
        refs = []
        for comp in components:
            if comp["type"] == "video" and not _is_blacklisted(comp["src"]):
                refs.append(ExternalRef(comp["display_name"], comp["src"], "video", folder))
            elif comp["type"] == "html":
                refs.extend(r for r in _refs_in_html(comp["content"], folder) if not _is_dead_footnote(r.url))
        return refs

    def _plan_linked(
        self, target: list[PlannedPage], pages: list[dict], folder: str, url_key: str
    ) -> None:
        for i, page in enumerate(pages, 1):
            path = PurePosixPath(folder) / f"{_safe_name(page['display_name'], i)}.html"
            planned = PlannedPage(page["display_name"], path, page["components"])
            self._by_url[_key(page[url_key])] = planned
            self._register(page[url_key], planned)
            target.append(planned)

    def _register(self, url: str | None, page: PlannedPage) -> None:
        if url:
            self._pages[_norm(url)] = page

    def _write_pages(self, root: Path, plan: Plan) -> None:
        for chapter in plan.chapters:
            self._write_page(root, chapter.index, self._chapter_body(chapter))
        for page in plan.content_pages():
            self._write_page(root, page, self._body(page) or _EMPTY_BODY)

    def _write_page(self, root: Path, page: PlannedPage, body: str) -> None:
        nav = f'<p><a href="{escape(_rel(_ROOT_INDEX, page.path))}">Back to index</a></p>\n'
        self._write_html(root / page.path, page.title, body, nav)

    @staticmethod
    def _write_html(target: Path, title: str, body: str, nav: str = "") -> None:
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text(
            _PAGE.format(title=escape(title), nav=nav, body=body), encoding="utf-8"
        )

    def _body(self, page: PlannedPage) -> str:
        return "\n".join(self._render(comp, page.path) for comp in page.components)

    def _chapter_body(self, chapter: PlannedChapter) -> str:
        index = chapter.index
        parts = [f"<h1>{escape(index.title)}</h1>"]
        if index.summary_html:
            parts.append(self._localise(index.summary_html, index.path))
        for title, verticals in chapter.sequentials:
            parts.append(f"<h2>{escape(title)}</h2>")
            parts.append(self._link_list(verticals, index.path))
        return "\n".join(parts)

    def _render(self, comp: dict, page: PurePosixPath, links_from: PurePosixPath | None = None) -> str:
        match comp["type"]:
            case "html":
                return self._localise(comp["content"], page, links_from)
            case "video":
                return self._video_html(comp)
            case "reading_link":
                return self._page_link(comp["reading_url"], page)
            case "dedup_link":
                return self._page_link(comp["dedup_url"], page)
            case other:
                raise ValueError(f"unknown component type: {other}")

    def _localise(self, html: str, page: PurePosixPath, links_from: PurePosixPath | None = None) -> str:
        prefix = _rel(PurePosixPath("assets"), page)
        html = STATIC_REF_RE.sub(
            lambda m: self._use(m.group(1), page, prefix), html
        )
        return self._strip_scripts(self._relink(html, links_from or page))

    def _target(self, url: str) -> PlannedPage | None:
        if page := self._pages.get(_norm(url)):
            return page
        segments = [s for s in urlsplit(url.strip()).path.split("/") if s]
        if segments[:1] == [_TEACHERS_SEGMENT] and len(segments) > 1:
            return self._teachers
        if segments[:1] == ["courses"] and len(segments) == 2:
            return self._home
        return None

    def _relink(self, html: str, base: PurePosixPath) -> str:
        """Point links at pages of the zip, and unwrap links to edX/DelftX/OCW hosts that have no copy in it."""
        soup = BeautifulSoup(html, "html.parser")
        media = {id(a) for a in soup.select(WIDGET_MEDIA_LINKS)}
        for link in soup.find_all("a", href=True):
            url = link["href"].strip()
            if id(link) in media or not _is_dead_host(urlsplit(url).hostname or ""):
                continue
            target = self._target(url)
            if target:
                link["href"] = _rel(target.path, base)
                continue
            if _is_dead_footnote(url):
                link["href"] = f"#{urlsplit(url).fragment}"
                continue
            if _norm(link.get_text()) != _norm(url):
                link.insert_after(NavigableString(f" ({url})"))
            link.unwrap()
        return str(soup)

    def _use(self, name: str, page: PurePosixPath, prefix: str) -> str:
        folder = _COURSE_OVERVIEW if page == _ROOT_INDEX else _folder(page)
        folders = self._usage.setdefault(asset_name(name), [])
        if folder not in folders:
            folders.append(folder)
        return f"{prefix}/{asset_name(name)}"

    def _page_link(self, url: str, page: PurePosixPath) -> str:
        target = self._by_url[_key(url)]
        return f'<p><a href="{escape(_rel(target.path, page))}">{escape(target.title)}</a></p>'

    @staticmethod
    def _video_link(comp: dict) -> str:
        return f'<p><a href="{escape(comp["src"])}">{escape(comp["display_name"])} (video)</a></p>'

    @classmethod
    def _video_html(cls, comp: dict) -> str:
        src, label = escape(comp["src"]), escape(comp["display_name"])
        return cls._video_link(comp) + f'<iframe src="{src}" title="{label}"></iframe>'

    def _strip_scripts(self, html: str) -> str:
        soup = BeautifulSoup(html, "html.parser")
        for tag in soup.find_all("script"):
            tag.decompose()
        for tag in soup.find_all(["img", "a"]):
            ref = tag.get("src") or tag.get("href")
            if isinstance(ref, str) and posixpath.basename(ref) in self._skipped_assets:
                tag.decompose()
        for tag in soup.find_all(True):
            for attr in [a for a in tag.attrs if a.lower().startswith("on")]:
                del tag[attr]
            href = tag.get("href")
            if isinstance(href, str) and href.strip().lower().startswith("javascript:"):
                del tag["href"]
        return str(soup)

    @staticmethod
    def _link_list(pages: list[PlannedPage], page: PurePosixPath) -> str:
        items = "".join(
            f'<li><a href="{escape(_rel(p.path, page))}">{escape(p.title)}</a></li>'
            for p in pages
        )
        return f"<ul>{items}</ul>"

    def _write_index(self, root: Path, plan: Plan) -> None:
        parts = [f"<h1>{escape(self.course.course_name)}</h1>"]
        parts += self._banner()
        if self.course.overview_summary_html:
            parts.append(self._localise(self.course.overview_summary_html, _ROOT_INDEX))
        for chapter in plan.chapters:
            parts.append(self._chapter_entry(chapter))
        for heading, pages in (("Readings", plan.readings), ("Shared lectures", plan.shared)):
            if pages:
                parts.append(f"<h2>{heading}</h2>")
                parts.append(self._link_list(pages, _ROOT_INDEX))
        parts.append(f'<p><a href="{_COURSE_DOCX}">Course overview (Docx)</a></p>')
        parts.append(f'<p><a href="{_EXTERNAL_XLSX}">External content, one row per occurrence (Excel)</a></p>')
        self._write_html(root / _ROOT_INDEX, self.course.course_name, "\n".join(parts))

    def _banner(self) -> list[str]:
        banner = self.course.banner_image_path
        names = [n for n, path in self.course.static_files.items() if path == banner and not is_skipped_asset(n)]
        return [f'<img src="{self._use(names[0], _ROOT_INDEX, "assets")}" alt="">'] if banner and names else []

    def _chapter_entry(self, chapter: PlannedChapter) -> str:
        head = (
            f'<h2><a href="{escape(_rel(chapter.index.path, _ROOT_INDEX))}">'
            f"{escape(chapter.index.title)}</a></h2>"
        )
        lists = "".join(
            f"<h3>{escape(title)}</h3>{self._link_list(verticals, _ROOT_INDEX)}"
            for title, verticals in chapter.sequentials
        )
        return head + lists

    def _copy_assets(self, root: Path) -> None:
        assets = root / "assets"
        assets.mkdir(parents=True, exist_ok=True)
        for name, source in self.course.static_files.items():
            if is_skipped_asset(name):
                self._skipped_assets.add(asset_name(name))
                log.info("Skipping blacklisted asset '%s'", name)
                continue
            shutil.copyfile(source, assets / asset_name(name))

    def _local_refs(self) -> list[ExternalRef]:
        refs = []
        for name in self.course.static_files:
            kind = _local_kind(name)
            if kind == "font":
                continue
            url = self.course.static_sources.get(name)
            if url is None:
                log.warning("No source url recorded for static file '%s', skipping its report row", name)
                continue
            skipped = is_skipped_asset(name)
            status = "left alone" if skipped else "copied"
            location = "" if skipped else f"assets/{asset_name(name)}"
            for folder in self._usage.get(asset_name(name), [_NOT_REFERENCED]):
                refs.append(ExternalRef(name, url, kind, folder, status, location))
        return refs

    @staticmethod
    def _zip(tmp: Path, out: Path) -> None:
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zf:
            for path in sorted(tmp.rglob("*")):
                if path.is_file():
                    zf.write(path, path.relative_to(tmp).as_posix())