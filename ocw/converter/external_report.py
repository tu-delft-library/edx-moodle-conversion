from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from urllib.parse import urlparse

from bs4 import BeautifulSoup
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ocw.utils import _XML_ILLEGAL_CHARS_RE

_OCW_HOST = "ocw.tudelft.nl"
# Domains never reported (each entry also covers its subdomains). Add boilerplate hosts here.
_BLACKLIST = frozenset({"creativecommons.org", "amazon.com"})
_HYPERLINK_REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"
_HTML_REFS = (("a", "href", "link"), ("iframe", "src", "embed"), ("img", "src", "image"))
_COLUMN_WIDTHS = (Inches(1.3), Inches(0.7), Inches(1.8), Inches(1.9), Inches(0.6), Inches(1.2))
_XLSX_HEADERS = ("File / link title", "Type", "URL", "Status", "Saved to", "Found at path")
_XLSX_COLUMN_WIDTHS = (32, 8, 60, 20, 36, 60)
_XLSX_COLUMN_COLORS = ("FFD9E1F2", "FFFCE4D6", "FF595959", "FFFFF2CC", "FFEAD1DC", "FFD9D2E9")
_XLSX_WRAP_COLUMNS = (3, 5, 6)
_XLSX_BORDER = Border(top=Side(style="thin", color="FF000000"), bottom=Side(style="thin", color="FF000000"))


@dataclass(frozen=True)
class ExternalRef:
    title: str
    url: str
    kind: str
    page: str
    status: str = "left alone"
    location: str = ""


_KIND_BY_SUFFIX = {
    ".pdf": "pdf",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".gif": "image",
    ".svg": "image",
    ".woff": "font",
    ".woff2": "font",
}


def _local_kind(name: str) -> str:
    return _KIND_BY_SUFFIX.get(PurePosixPath(name).suffix.lower(), "file")


def _in_domain(host: str, domain: str) -> bool:
    return host == domain or host.endswith(f".{domain}")


def _is_blacklisted(url: str) -> bool:
    host = urlparse(url).hostname or ""
    return any(_in_domain(host, domain) for domain in _BLACKLIST)


def _is_external(url: str) -> bool:
    parsed = urlparse(url)
    host = parsed.hostname or ""
    return (
        parsed.scheme in ("http", "https")
        and not _in_domain(host, _OCW_HOST)
        and not _is_blacklisted(url)
    )


def _fallback_title(url: str) -> str:
    parsed = urlparse(url)
    last = parsed.path.rstrip("/").rsplit("/", 1)[-1]
    return f"{parsed.hostname}/{last}" if last else parsed.hostname


def _refs_in_html(html: str, page: str) -> list[ExternalRef]:
    soup = BeautifulSoup(html, "lxml")
    refs = []
    for tag, attr, kind in _HTML_REFS:
        for el in soup.find_all(tag, **{attr: True}):
            url = el[attr]
            if _is_external(url):
                title = el.get("title") or el.get("alt") or el.get_text(" ", strip=True)
                refs.append(ExternalRef(title.strip() or _fallback_title(url), url, kind, page))
    return refs


def _clean(text: str) -> str:
    return _XML_ILLEGAL_CHARS_RE.sub("", text)


def merge_refs(refs: list[ExternalRef]) -> list[tuple[ExternalRef, list[str]]]:
    """Collapse refs with an identical URL into one entry, keeping every page they were found on."""
    merged: dict[str, tuple[ExternalRef, list[str]]] = {}
    for ref in refs:
        _, pages = merged.setdefault(ref.url, (ref, []))
        if ref.page not in pages:
            pages.append(ref.page)
    return list(merged.values())


def _found_on(pages: list[str]) -> str:
    if len(pages) == 1:
        return pages[0]
    return f"{len(pages)} locations ({pages[0]})"


class ExternalContentReport:
    """Write external references as a Docx table with clickable URLs."""

    def __init__(self, refs: list[ExternalRef]) -> None:
        self.refs = refs

    def write(self, path: Path) -> None:
        doc = Document()
        section = doc.sections[0]
        section.left_margin = section.right_margin = Inches(0.5)
        section.top_margin = section.bottom_margin = Inches(0.5)
        doc.styles["Normal"].font.size = Pt(9)
        doc.add_heading("External content", level=0)
        table = doc.add_table(rows=1, cols=6)
        table.autofit = False
        for cell, label in zip(table.rows[0].cells, ("Title", "Type", "URL", "Found on", "Status", "Saved to")):
            cell.text = label
            cell.paragraphs[0].runs[0].bold = True
        merged = sorted(merge_refs(self.refs), key=lambda rp: rp[0].location == "")
        for ref, pages in merged:
            row = table.add_row().cells
            row[0].text, row[1].text = _clean(ref.title), ref.kind
            row[3].text, row[4].text = _clean(_found_on(pages)), ref.status
            row[5].text = ref.location
            self._add_hyperlink(row[2].paragraphs[0], ref.url)
        for i, width in enumerate(_COLUMN_WIDTHS):
            table.columns[i].width = width
            for row in table.rows:
                row.cells[i].width = width
        doc.save(path)

    def write_xlsx(self, path: Path) -> None:
        """Write one row per (file, path it was found on) observation, unlike the merged Docx table."""
        wb = Workbook()
        ws = wb.active
        ws.title = "External content"
        ws.append(_XLSX_HEADERS)
        header_fill = PatternFill(start_color="FF4472C4", end_color="FF4472C4", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
            cell.border = _XLSX_BORDER
        refs = sorted(self.refs, key=lambda ref: ref.location == "")
        for ref in refs:
            ws.append((ref.title, ref.kind, ref.url, ref.status, ref.location or "-", ref.page))
        for col_idx, color in enumerate(_XLSX_COLUMN_COLORS, start=1):
            fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            for row_idx in range(2, ws.max_row + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill
        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=3)
            cell.hyperlink = cell.value
            cell.font = Font(color="FF9CC3E5", underline="single")
        for row_idx in range(2, ws.max_row + 1):
            for col_idx in range(1, len(_XLSX_HEADERS) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = _XLSX_BORDER
                if col_idx in _XLSX_WRAP_COLUMNS:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col_idx, width in enumerate(_XLSX_COLUMN_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.freeze_panes = "A2"
        wb.save(path)

    @staticmethod
    def _add_hyperlink(paragraph, url: str) -> None:
        rel_id = paragraph.part.relate_to(url, _HYPERLINK_REL, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel_id)
        run, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "0563C1")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "single")
        rpr.append(color)
        rpr.append(underline)
        text = OxmlElement("w:t")
        text.text = _clean(url)
        run.append(rpr)
        run.append(text)
        link.append(run)
        paragraph._p.append(link)