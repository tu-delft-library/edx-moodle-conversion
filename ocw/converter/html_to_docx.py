import logging
import posixpath
import re
from dataclasses import dataclass, replace
from pathlib import Path
from urllib.parse import urlsplit

from bs4 import BeautifulSoup, NavigableString, Tag
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, UnidentifiedImageError

from ocw.converter.external_report import _HYPERLINK_REL
from ocw.utils import _XML_ILLEGAL_CHARS_RE

log = logging.getLogger("ocw.html_to_docx")

_SPACE_RE = re.compile(r"\s+")
_LEADING_RELATIVE_RE = re.compile(r"^(?:\.{1,2}/)+")
_SPAN_RE = re.compile(r"\d+")
_HEADINGS = frozenset({"h1", "h2", "h3", "h4", "h5", "h6"})
_BLOCK_TAGS = frozenset(
    {"p", "div", "section", "article", "figure", "figcaption", "blockquote", "ul", "ol", "table", "details",
     "summary", "pre", "li", *_HEADINGS}
)
_SKIPPED_TAGS = frozenset({"script", "style", "noscript", "svg", "hr", "head", "meta", "link"})
_KNOWN_INLINE = frozenset({"span", "u", "small", "mark", "abbr", "cite", "label", "font", "s", "strike", "del", "ins"})
_NATIVE_IMAGE_FORMATS = frozenset({"PNG", "JPEG", "GIF", "BMP", "TIFF"})
_FIXED_HEADING_ATTR = "data-structural"
_PAGE_BREAK_ATTR = "data-page-break"
_DOWNLOAD_WIDGET_CLASS = "TUD_box_download"
_DOWNLOAD_GROUPS = (("video-download", "Video"), ("subtitle-download", "Subtitles"), ("other-download", "Other"))
_DOWNLOAD_HEADER = ("Type", "File", "Link")
_MAX_HEADING_LEVEL = 9
_MAX_LIST_LEVEL = 3
_MONO_FONT = "Courier New"
_LINK_COLOR = "0563C1"


@dataclass(frozen=True)
class _Format:
    bold: bool = False
    italic: bool = False
    mono: bool = False
    superscript: bool = False
    subscript: bool = False


class _Para:
    """A paragraph in `target` that is only created once something is written to it."""

    def __init__(self, target, style: str | None = None) -> None:
        self.target = target
        self.style = style
        self._paragraph = None

    @property
    def started(self) -> bool:
        return self._paragraph is not None

    def get(self):
        if self._paragraph is None:
            self._paragraph = self.target.add_paragraph(style=self.style)
        return self._paragraph

    def close(self) -> None:
        """Trim trailing whitespace left by the source markup."""
        if self._paragraph is None:
            return
        runs = [r for r in self._paragraph.runs if r.text]
        if runs:
            runs[-1].text = runs[-1].text.rstrip()


def _clean(text: str) -> str:
    return _XML_ILLEGAL_CHARS_RE.sub("", text)


def _span(cell: Tag, attr: str) -> int:
    match = _SPAN_RE.match(str(cell.get(attr) or ""))
    return max(int(match.group()), 1) if match else 1


def _local_url(href: str) -> str:
    """Drop the page relative `../` prefix: the docx sits at the zip root next to `assets/`."""
    parts = urlsplit(href)
    if parts.scheme or href.startswith("//"):
        return href
    return _LEADING_RELATIVE_RE.sub("", href)


class HtmlToDocx:
    """Write course HTML into an existing python-docx document, with minimal styling."""

    def __init__(self, doc, assets: Path, scratch: Path) -> None:
        self.doc = doc
        self.assets = assets
        self.scratch = scratch
        self._heading_offset = 0
        self._unhandled: set[str] = set()
        section = doc.sections[0]
        self._usable_width = section.page_width - section.left_margin - section.right_margin

    def add(self, html: str, heading_offset: int = 0) -> None:
        """Append `html` to the document, shifting every heading level down by `heading_offset`.

        Headings carrying a `data-structural` attribute keep their level, a `data-page-break` heading starts a page.
        """
        self._heading_offset = heading_offset
        self._blocks(BeautifulSoup(html, "html.parser"), self.doc)

    def _blocks(self, parent, target) -> None:
        para = _Para(target)
        for child in parent.children:
            if isinstance(child, Tag) and child.name in _BLOCK_TAGS:
                para.close()
                self._block(child, target)
                para = _Para(target)
            else:
                self._inline(child, para, _Format())
        para.close()

    def _block(self, node: Tag, target) -> None:
        name = node.name
        if name in _HEADINGS:
            offset = 0 if node.has_attr(_FIXED_HEADING_ATTR) else self._heading_offset
            level = min(int(name[1]) + offset, _MAX_HEADING_LEVEL)
            para = _Para(target, f"Heading {level}")
            self._write_block(node, para, _Format())
            if node.has_attr(_PAGE_BREAK_ATTR) and para.started:
                para.get().paragraph_format.page_break_before = True
        elif name in ("ul", "ol"):
            self._list(node, target, 0)
        elif name == "table":
            self._table(node, target)
        elif name == "summary":
            self._write_block(node, _Para(target), _Format(bold=True))
        elif name == "pre":
            self._pre(node, target)
        elif _DOWNLOAD_WIDGET_CLASS in (node.get("class") or []):
            self._download_table(node, target)
        else:
            self._blocks(node, target)

    def _write_block(self, node: Tag, para: _Para, fmt: _Format) -> None:
        self._inline_children(node, para, fmt)
        para.close()

    def _inline_children(self, node: Tag, para: _Para, fmt: _Format) -> None:
        for child in node.children:
            self._inline(child, para, fmt)

    def _inline(self, node, para: _Para, fmt: _Format) -> None:
        if isinstance(node, NavigableString):
            if type(node) is NavigableString:
                self._text(para, str(node), fmt)
            return
        name = node.name
        if name in _SKIPPED_TAGS:
            log.debug("Skipping <%s> element", name)
        elif name == "br":
            para.get().add_run().add_break()
        elif name in ("strong", "b"):
            self._inline_children(node, para, replace(fmt, bold=True))
        elif name in ("em", "i"):
            self._inline_children(node, para, replace(fmt, italic=True))
        elif name in ("code", "kbd", "samp", "tt"):
            self._inline_children(node, para, replace(fmt, mono=True))
        elif name == "sup":
            self._inline_children(node, para, replace(fmt, superscript=True))
        elif name == "sub":
            self._inline_children(node, para, replace(fmt, subscript=True))
        elif name == "a":
            self._link(node, para, fmt)
        elif name == "img":
            self._image(node, para)
        elif name == "iframe":
            self._embed(node, para, fmt)
        else:
            self._nested(node, para, fmt)

    def _nested(self, node: Tag, para: _Para, fmt: _Format) -> None:
        """Flatten a container met inside inline content (a `p` in an `li`, a stray unknown tag)."""
        if node.name in _BLOCK_TAGS:
            if para.started:
                para.get().add_run().add_break()
        elif node.name not in _KNOWN_INLINE and node.name not in self._unhandled:
            self._unhandled.add(node.name)
            log.info("Unhandled <%s> element, writing its children as plain text", node.name)
        self._inline_children(node, para, fmt)

    def _text(self, para: _Para, text: str, fmt: _Format) -> None:
        text = _SPACE_RE.sub(" ", _clean(text))
        if not para.started:
            text = text.lstrip()
        if text:
            self._style(para.get().add_run(text), fmt)

    @staticmethod
    def _style(run, fmt: _Format) -> None:
        run.bold = fmt.bold or None
        run.italic = fmt.italic or None
        if fmt.superscript:
            run.font.superscript = True
        if fmt.subscript:
            run.font.subscript = True
        if fmt.mono:
            run.font.name = _MONO_FONT

    def _link(self, node: Tag, para: _Para, fmt: _Format) -> None:
        href = str(node.get("href") or "").strip()
        if not href or href.startswith("#") or node.find("img"):
            self._inline_children(node, para, fmt)
            return
        text = _SPACE_RE.sub(" ", node.get_text()).strip() or href
        self._hyperlink(para, _local_url(href), text, fmt)

    def _embed(self, node: Tag, para: _Para, fmt: _Format) -> None:
        src = str(node.get("src") or "").strip()
        if src:
            self._hyperlink(para, src, src, fmt)

    def _hyperlink(self, para: _Para, url: str, text: str, fmt: _Format) -> None:
        paragraph = para.get()
        rel_id = paragraph.part.relate_to(url, _HYPERLINK_REL, is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rel_id)
        run, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
        for tag, value in (("w:b", fmt.bold), ("w:i", fmt.italic)):
            if value:
                rpr.append(OxmlElement(tag))
        color, underline = OxmlElement("w:color"), OxmlElement("w:u")
        color.set(qn("w:val"), _LINK_COLOR)
        underline.set(qn("w:val"), "single")
        rpr.append(color)
        rpr.append(underline)
        label = OxmlElement("w:t")
        label.text = _clean(text)
        label.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        run.append(rpr)
        run.append(label)
        link.append(run)
        paragraph._p.append(link)

    def _image(self, node: Tag, para: _Para) -> None:
        src = str(node.get("src") or "").strip()
        name = posixpath.basename(urlsplit(src).path)
        is_asset = "assets/" in src
        local = is_asset and (self.assets / name).is_file()
        if local and name.lower().endswith(".svg"):
            log.debug("Not embedding svg '%s', it stays in assets/", name)
            return
        prepared = self._prepare_image(self.assets / name) if local else None
        if prepared is None:
            if not local:
                log.warning("Image '%s' is not a local asset, writing a link instead", src)
            href = _local_url(src) if is_asset else src
            self._hyperlink(para, href, f"figure: {name or src}", _Format(italic=True))
            return
        shape = para.get().add_run().add_picture(str(prepared))
        if shape.width > self._usable_width:
            shape.height = int(shape.height * self._usable_width / shape.width)
            shape.width = self._usable_width

    def _prepare_image(self, path: Path) -> Path | None:
        """The file itself when Word can embed it, a png copy for other formats, `None` when unreadable."""
        try:
            with Image.open(path) as img:
                if img.format in _NATIVE_IMAGE_FORMATS:
                    return path
                converted = self.scratch / f"{path.name}.png"
                img.convert("RGBA").save(converted)
                return converted
        except (UnidentifiedImageError, OSError) as exc:
            log.warning("Could not read image '%s': %s", path.name, exc)
            return None

    def _list(self, node: Tag, target, depth: int) -> None:
        base = "List Number" if node.name == "ol" else "List Bullet"
        level = min(depth + 1, _MAX_LIST_LEVEL)
        style = base if level == 1 else f"{base} {level}"
        for item in node.find_all("li", recursive=False):
            para = _Para(target, style)
            for child in item.children:
                if isinstance(child, Tag) and child.name in ("ul", "ol"):
                    self._list(child, target, depth + 1)
                else:
                    self._inline(child, para, _Format())
            para.close()

    def _table(self, node: Tag, target) -> None:
        rows = [
            row.find_all(["td", "th"], recursive=False)
            for row in node.find_all("tr")
            if row.find_parent("table") is node
        ]
        cols = max((sum(_span(c, "colspan") for c in row) for row in rows), default=0)
        if not cols:
            return
        table = target.add_table(rows=len(rows), cols=cols)
        table.style = "Table Grid"
        for r, row in enumerate(rows):
            col = 0
            for cell_node in row:
                span = min(_span(cell_node, "colspan"), cols - col)
                cell = table.cell(r, col)
                if span > 1:
                    cell = cell.merge(table.cell(r, col + span - 1))
                self._blocks(cell_node, cell)
                self._tidy_cell(cell)
                col += span

    @staticmethod
    def _tidy_cell(cell) -> None:
        first = cell.paragraphs[0]
        if len(cell.paragraphs) > 1 and not first._p.xpath("./w:r | ./w:hyperlink"):
            first._p.getparent().remove(first._p)

    def _download_table(self, node: Tag, target) -> None:
        """The DelftxDownloads widget as a Type / File / Link table; nothing when it has no real links."""
        rows = []
        for group_class, group in _DOWNLOAD_GROUPS:
            for link in node.select(f".{group_class} li a[href]"):
                href = str(link["href"]).strip()
                if href and not href.startswith("#"):
                    rows.append((group, _SPACE_RE.sub(" ", link.get_text(" ")).strip(), _local_url(href)))
        if not rows:
            log.warning("Download widget has no real links, dropping it (tool page probably not fetched)")
            return
        table = target.add_table(rows=len(rows) + 1, cols=len(_DOWNLOAD_HEADER))
        table.style = "Table Grid"
        for cell, title in zip(table.rows[0].cells, _DOWNLOAD_HEADER):
            cell.paragraphs[0].add_run(title).bold = True
        for r, (group, label, href) in enumerate(rows, 1):
            cells = table.rows[r].cells
            cells[0].paragraphs[0].add_run(group)
            cells[1].paragraphs[0].add_run(_clean(label))
            self._hyperlink(_Para(cells[2]), href, href, _Format())
            self._tidy_cell(cells[2])

    def _pre(self, node: Tag, target) -> None:
        para = _Para(target)
        for i, line in enumerate(_clean(node.get_text()).split("\n")):
            if i:
                para.get().add_run().add_break()
            if line:
                self._style(para.get().add_run(line), _Format(mono=True))
