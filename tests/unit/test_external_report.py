import zipfile

from docx import Document

from ocw.converter.external_report import (
    ExternalContentReport,
    ExternalRef,
    _fallback_title,
    _found_on,
    _is_blacklisted,
    _is_external,
    _refs_in_html,
    collect_external,
    merge_refs,
)
from ocw.converter import external_report
from ocw.converter.zip_builder import ZipBuilder
from tests.unit.test_zip_builder import _site


def _rows(path) -> list[list[str]]:
    table = Document(path).tables[0]
    return [[c.text for c in row.cells] for row in table.rows]


def test_is_external_excludes_ocw_hosts_and_non_http():
    assert _is_external("https://example.com/x")
    assert not _is_external("https://ocw.tudelft.nl/x")
    assert not _is_external("https://cdn.ocw.tudelft.nl/x")
    assert not _is_external("mailto:a@b.nl")
    assert not _is_external("/relative/path")
    assert _is_external("https://notocw.tudelft.nl.evil.com/x")


def test_fallback_title_uses_host_and_last_segment():
    assert _fallback_title("https://example.com/a/b/paper.pdf") == "example.com/paper.pdf"
    assert _fallback_title("https://example.com/") == "example.com"


def test_refs_in_html_prefers_title_then_alt_then_text_then_fallback():
    html = (
        '<a href="https://a.com/1" title="Titled">Text</a>'
        '<a href="https://a.com/2">Plain text</a>'
        '<img src="https://a.com/3.png" alt="Alt text">'
        '<iframe src="https://a.com/embed/4"></iframe>'
        '<a href="https://ocw.tudelft.nl/skip">Skip</a>'
    )
    refs = _refs_in_html(html, "Page")
    assert [(r.title, r.kind) for r in refs] == [
        ("Titled", "link"),
        ("Plain text", "link"),
        ("a.com/4", "embed"),
        ("Alt text", "image"),
    ]


def test_collect_external_lists_videos_and_links_from_all_pages():
    course = _site().course(keep_all_content=True)
    course.parse()
    refs = collect_external(course)
    videos = [r for r in refs if r.kind == "video"]
    assert [(v.title, v.url, v.page) for v in videos] == [
        ("Lecture 1", "https://www.youtube.com/embed/abc123", "Lecture 1")
    ]


def test_report_round_trips_through_docx(tmp_path):
    refs = [
        ExternalRef("Intro video", "https://youtu.be/x", "video", "Lecture 1"),
        ExternalRef("Paper", "https://example.com/p.pdf", "link", "Reading 1"),
    ]
    path = tmp_path / "report.docx"
    ExternalContentReport(refs).write(path)
    assert _rows(path) == [
        ["Title", "Type", "URL", "Found on"],
        ["Intro video", "video", "https://youtu.be/x", "Lecture 1"],
        ["Paper", "link", "https://example.com/p.pdf", "Reading 1"],
    ]


def test_identical_urls_merge_into_one_entry_with_all_pages():
    refs = [
        ExternalRef("Cress", "http://www.cress.nl/About.aspx", "link", "Exams"),
        ExternalRef("Cress", "http://www.cress.nl/About.aspx", "link", "Software"),
        ExternalRef("Cress", "http://www.cress.nl/About.aspx", "link", "Software"),
    ]
    merged = merge_refs(refs)
    assert len(merged) == 1
    ref, pages = merged[0]
    assert ref.url == "http://www.cress.nl/About.aspx"
    assert pages == ["Exams", "Software"]


def test_found_on_wording_by_page_count():
    assert _found_on(["Exams"]) == "Exams"
    assert _found_on(["Exams", "Software"]) == "found on both: Exams; Software"
    assert _found_on(["A", "B", "C"]) == "found on 3 pages: A; B; C"


def test_report_urls_are_clickable_hyperlinks(tmp_path):
    path = tmp_path / "report.docx"
    ExternalContentReport([ExternalRef("T", "https://example.com/a", "link", "P")]).write(path)
    targets = [rel.target_ref for rel in Document(path).part.rels.values() if rel.is_external]
    assert targets == ["https://example.com/a"]


def test_report_strips_xml_illegal_characters(tmp_path):
    path = tmp_path / "report.docx"
    ExternalContentReport([ExternalRef("bad\x0btitle", "https://e.com/", "link", "P")]).write(path)
    assert _rows(path)[1][0] == "badtitle"


def test_empty_report_is_header_only(tmp_path):
    path = tmp_path / "report.docx"
    ExternalContentReport([]).write(path)
    assert len(_rows(path)) == 1


def test_zip_contains_docx_at_root_and_index_links_it(tmp_path):
    course = _site().course(keep_all_content=True)
    course.parse()
    out = tmp_path / "course.zip"
    ZipBuilder(course).build(out)
    with zipfile.ZipFile(out) as zf:
        assert "Example_Course/external_content.docx" in zf.namelist()
        assert 'href="external_content.docx"' in zf.read("Example_Course/index.html").decode()
        zf.extract("Example_Course/external_content.docx", tmp_path)
    urls = [row[2] for row in _rows(tmp_path / "Example_Course/external_content.docx")]
    assert "https://www.youtube.com/embed/abc123" in urls


CC_FOOTER = (
    '<section class="license"><a href="http://creativecommons.org/licenses/by-nc-sa/4.0/" rel="license">'
    '<img alt="Creative Commons License" src="https://i.creativecommons.org/l/by-nc-sa/4.0/88x31.png"></a>'
    '</section>'
)


def test_blacklist_covers_domain_and_subdomains_only():
    assert _is_blacklisted("http://creativecommons.org/licenses/by/4.0/")
    assert _is_blacklisted("https://i.creativecommons.org/l/by/4.0/88x31.png")
    assert not _is_blacklisted("https://notcreativecommons.org/x")
    assert not _is_blacklisted("https://example.com/creativecommons.org")


def test_blacklisted_hosts_are_not_reported():
    assert _refs_in_html(CC_FOOTER, "Page") == []
    assert not _is_external("https://i.creativecommons.org/l/by/4.0/88x31.png")


def test_blacklist_entries_are_configurable(monkeypatch):
    monkeypatch.setattr(external_report, "_BLACKLIST", frozenset({"example.com"}))
    assert _refs_in_html('<a href="https://sub.example.com/x">x</a>', "Page") == []
    assert len(_refs_in_html(CC_FOOTER, "Page")) == 2


def test_amazon_links_are_not_reported():
    html = '<a href="https://www.amazon.com/dp/123">Buy the book</a><a href="https://amazon.com/x">Again</a>'
    assert _refs_in_html(html, "Page") == []


def test_blacklisted_video_is_not_reported():
    class _Course:
        overview_summary_html = None
        chapters = [
            {
                "display_name": "Ch",
                "sequentials": [
                    {
                        "display_name": "S",
                        "verticals": [
                            {
                                "display_name": "V",
                                "components": [
                                    {"type": "video", "display_name": "A", "src": "https://creativecommons.org/v"},
                                    {"type": "video", "display_name": "B", "src": "https://youtu.be/x"},
                                ],
                            }
                        ],
                    }
                ],
            }
        ]
        reading_pages: list = []
        dedup_pages: list = []

    assert [r.title for r in collect_external(_Course())] == ["B"]
