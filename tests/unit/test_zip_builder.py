import io
import logging
import zipfile
from pathlib import PurePosixPath

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from PIL import Image

from ocw.converter.zip_builder import (
    ZipBuilder,
    _EMPTY_BODY,
    _key,
    _safe_name,
)
from tests.wp_builders import WPFixtureSite

LECTURE_URL = "https://ocw.tudelft.nl/course-lectures/lec1/"
EXERCISE_URL = "https://ocw.tudelft.nl/course-exercises/ex1/"
READING_URL = "https://ocw.tudelft.nl/course-readings/read1/"
SUBJECT_URL = "https://ocw.tudelft.nl/courses/example/subjects/1-intro/"


def _site(empty_lecture: bool = False) -> WPFixtureSite:
    site = WPFixtureSite()
    site.add_home([(SUBJECT_URL, "1. Intro")])
    site.add_subject_page(
        SUBJECT_URL,
        f"""
        <ul class="activities">
        <li>
        <h4><span class="course-title">1.1 Basics</span></h4>
        <ul>
        <li><a class="icon icon--lecture" href="{LECTURE_URL}">Lecture 1</a></li>
        <li><a class="icon icon--exercise" href="{EXERCISE_URL}">Exercise 1</a></li>
        <li><a class="icon icon--reading" href="{READING_URL}">Reading 1</a></li>
        </ul>
        </li>
        </ul>
        """,
        intro_html='<div class="vc_row"><p>Subject intro.</p></div>',
    )
    lecture_html = (
        '<div class="vc_row"></div>'
        if empty_lecture
        else """
        <h1>Lecture 1</h1>
        <div class="vc_row"><p onclick="steal()">Text.</p><script>alert(1)</script></div>
        <div class="vc_row"><img src="https://ocw.tudelft.nl/wp-content/uploads/diagram.png"></div>
        <div class="vc_row"><img src="https://ocw.tudelft.nl/wp-content/uploads/88x31.png"></div>
        <div class="vc_row"><iframe src="https://www.youtube.com/embed/abc123"></iframe></div>
        """
    )
    site.add_lecture_page(LECTURE_URL, lecture_html)
    site.add_lecture_page(EXERCISE_URL, "<h1>Exercise 1</h1><p>Solve it.</p>")
    site.add_reading_page(
        READING_URL,
        '<div class="vc_download"><a class="icon fa-file" '
        'href="https://ocw.tudelft.nl/wp-content/uploads/Chapter.pdf">'
        "<strong>Chapter</strong></a></div>",
        pdf_name="Chapter.pdf",
    )
    site.add_static_file("diagram.png", b"png-bytes")
    site.add_static_file("88x31.png", b"badge")
    site.add_static_file("fontawesome-webfont.woff2", b"font")
    site.add_static_file("fontawesome-webfont.woff", b"font")
    return site


@pytest.fixture
def archive(tmp_path):
    course = _site().course(keep_all_content=True)
    course.parse()
    out = tmp_path / "course.zip"
    ZipBuilder(course).build(out)
    with zipfile.ZipFile(out) as zf:
        yield zf


def _read(zf: zipfile.ZipFile, name: str) -> str:
    return zf.read(name).decode("utf-8")


def test_zip_has_folder_tree_under_course_slug(archive):
    names = set(archive.namelist())
    assert "Example_Course/index.html" in names
    assert "Example_Course/01_1._Intro/index.html" in names
    assert "Example_Course/01_1._Intro/01_1.1_Basics/01_Lecture_1.html" in names
    assert "Example_Course/01_1._Intro/01_1.1_Basics/02_Exercise_1.html" in names


def test_assets_are_flat_under_assets(archive):
    names = set(archive.namelist())
    assert "Example_Course/assets/diagram.png" in names
    assert "Example_Course/assets/Chapter.pdf" in names
    assert archive.read("Example_Course/assets/diagram.png") == b"png-bytes"


def test_static_references_point_at_relative_assets(archive):
    page = _read(archive, "Example_Course/01_1._Intro/01_1.1_Basics/01_Lecture_1.html")
    assert 'src="../../assets/diagram.png"' in page
    assert "/static/" not in page


def test_scripts_and_event_handlers_are_stripped(archive):
    page = _read(archive, "Example_Course/01_1._Intro/01_1.1_Basics/01_Lecture_1.html")
    assert "<script" not in page
    assert "onclick" not in page
    assert "Text." in page


def test_kept_title_and_video_embed_are_in_the_page(archive):
    page = _read(archive, "Example_Course/01_1._Intro/01_1.1_Basics/01_Lecture_1.html")
    assert "<h1>Lecture 1</h1>" in page
    assert 'src="https://www.youtube.com/embed/abc123"' in page


def test_reading_link_resolves_to_reading_page(archive):
    page = _read(archive, "Example_Course/01_1._Intro/01_1.1_Basics/03_Reading_1.html")
    assert 'href="../../readings/01_Reading_1.html"' in page
    assert "Example_Course/readings/01_Reading_1.html" in archive.namelist()


def test_chapter_index_lists_summary_and_pages(archive):
    page = _read(archive, "Example_Course/01_1._Intro/index.html")
    assert "<h1>1. Intro</h1>" in page
    assert "Subject intro." in page
    assert 'href="01_1.1_Basics/01_Lecture_1.html"' in page


def test_root_index_lists_chapters_and_readings(archive):
    page = _read(archive, "Example_Course/index.html")
    assert "<h1>Example Course</h1>" in page
    assert 'href="01_1._Intro/index.html"' in page
    assert 'href="readings/01_Reading_1.html"' in page
    assert "Back to index" not in page


def test_pages_link_back_to_root_index(archive):
    page = _read(archive, "Example_Course/01_1._Intro/01_1.1_Basics/01_Lecture_1.html")
    assert 'href="../../index.html">Back to index' in page


def test_dedup_link_resolves_to_shared_page(tmp_path):
    site = _site()
    second = "https://ocw.tudelft.nl/courses/example/subjects/2-more/"
    site.add_home([(SUBJECT_URL, "1. Intro"), (second, "2. More")])
    site.add_subject_page(
        second,
        f"""
        <ul class="activities"><li>
        <h4><span class="course-title">2.1 Again</span></h4>
        <ul><li><a class="icon icon--lecture" href="{LECTURE_URL}">Lecture 1</a></li></ul>
        </li></ul>
        """,
    )
    course = site.course(keep_all_content=True)
    course.parse()
    out = tmp_path / "shared.zip"
    ZipBuilder(course).build(out)
    with zipfile.ZipFile(out) as zf:
        assert "Example_Course/shared/01_Lecture_1.html" in zf.namelist()
        page = _read(zf, "Example_Course/01_1._Intro/01_1.1_Basics/01_Lecture_1.html")
        assert 'href="../../shared/01_Lecture_1.html"' in page


def test_unknown_component_type_raises():
    builder = ZipBuilder(_site().course())
    with pytest.raises(ValueError, match="unknown component type: nope"):
        builder._render({"type": "nope"}, PurePosixPath("a.html"))


def test_safe_name_neutralises_path_traversal():
    assert _safe_name("../../etc/passwd", 3) == "03_etc_passwd"


def test_safe_name_falls_back_for_empty_titles():
    assert _safe_name("///", 1) == "01_untitled"


def test_key_drops_fragment_and_trailing_slash():
    assert _key("https://x/y/#frag") == "https://x/y"


def test_skipped_asset_is_logged_at_info(tmp_path, caplog):
    course = _site().course(keep_all_content=True)
    course.parse()
    with caplog.at_level(logging.INFO, logger="ocw.zip_builder"):
        ZipBuilder(course).build(tmp_path / "course.zip")
    assert "Skipping blacklisted asset '88x31.png'" in caplog.text


def test_empty_vertical_writes_placeholder(tmp_path):
    course = _site(empty_lecture=True).course(keep_all_content=True)
    course.parse()
    out = tmp_path / "course.zip"
    ZipBuilder(course).build(out)
    with zipfile.ZipFile(out) as zf:
        pages = [zf.read(n).decode("utf-8") for n in zf.namelist() if n.endswith(".html")]
    assert any(_EMPTY_BODY in page for page in pages)


def test_usage_recorded_per_page(tmp_path):
    course = _site().course(keep_all_content=True)
    course.parse()
    builder = ZipBuilder(course)
    builder.build(tmp_path / "course.zip")
    assert builder._usage["diagram.png"] == ["01_1._Intro/01_1.1_Basics/"]


def test_local_refs_location_set_for_copied_file(tmp_path):
    course = _site().course(keep_all_content=True)
    course.parse()
    builder = ZipBuilder(course)
    builder.build(tmp_path / "course.zip")
    ref = next(r for r in builder._local_refs() if r.title == "diagram.png")
    assert ref.location == "assets/diagram.png"
    assert ref.status == "copied"


def test_local_refs_location_blank_for_skipped_file(tmp_path):
    course = _site().course(keep_all_content=True)
    course.parse()
    builder = ZipBuilder(course)
    builder.build(tmp_path / "course.zip")
    ref = next(r for r in builder._local_refs() if r.title == "88x31.png")
    assert ref.location == ""
    assert ref.status == "left alone"


def test_local_refs_labels_assets_not_used_in_any_page_body(tmp_path):
    banner_url = "https://ocw.tudelft.nl/wp-content/uploads/banner.jpg"
    site = _site()
    site.add_home([(SUBJECT_URL, "1. Intro")], banner_url=banner_url)
    site.add_static_file("banner.jpg", b"banner-bytes")
    course = site.course(keep_all_content=True)
    course.parse()
    builder = ZipBuilder(course)
    builder.build(tmp_path / "course.zip")
    ref = next(r for r in builder._local_refs() if r.title == "banner.jpg")
    assert ref.page == "not referenced on any page"


def test_local_refs_drops_font_rows(tmp_path):
    course = _site().course(keep_all_content=True)
    course.parse()
    builder = ZipBuilder(course)
    builder.build(tmp_path / "course.zip")
    assert not any(r.kind == "font" for r in builder._local_refs())
    assert not any("fontawesome" in r.title for r in builder._local_refs())


def test_empty_vertical_keeps_chapter_index_intact(tmp_path):
    course = _site(empty_lecture=True).course(keep_all_content=True)
    course.parse()
    out = tmp_path / "course.zip"
    ZipBuilder(course).build(out)
    with zipfile.ZipFile(out) as zf:
        index = _read(zf, "Example_Course/01_1._Intro/index.html")
        lecture = _read(zf, "Example_Course/01_1._Intro/01_1.1_Basics/01_Lecture_1.html")
    assert "Subject intro." in index
    assert 'href="01_1.1_Basics/01_Lecture_1.html"' in index
    assert _EMPTY_BODY in lecture
    assert _EMPTY_BODY not in index


@pytest.mark.parametrize("suffix", ["#page=3", "?download=1", "?a=1#b"])
def test_static_ref_keeps_fragment_and_query_off_the_asset_name(suffix):
    builder = ZipBuilder(_site().course())
    html = builder._localise(f'<a href="/static/Chapter.pdf{suffix}">c</a>', PurePosixPath("a/b.html"))
    assert f'href="../assets/Chapter.pdf{suffix}"' in html
    assert builder._usage == {"Chapter.pdf": ["a/"]}


def _png_bytes() -> bytes:
    buffer = io.BytesIO()
    Image.new("RGB", (40, 30), "red").save(buffer, format="PNG")
    return buffer.getvalue()


def _docx(tmp_path, site=None) -> Document:
    course = (site or _site()).course(keep_all_content=True)
    course.parse()
    out = tmp_path / "course.zip"
    ZipBuilder(course).build(out)
    with zipfile.ZipFile(out) as zf:
        return Document(io.BytesIO(zf.read("Example_Course/course.docx")))


def _headings(doc) -> list[tuple[str, str]]:
    return [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]


def test_docx_starts_with_the_external_table_then_contents_then_the_course(tmp_path):
    doc = _docx(tmp_path)
    blocks = ["<table>" if isinstance(b, Table) else b.text for b in doc.iter_inner_content()]
    assert blocks[0] == "External content"
    assert blocks.index("<table>") < blocks.index("Contents") < blocks.index("Example Course")


def test_docx_contents_lists_chapters_and_sequentials(tmp_path):
    doc = _docx(tmp_path)
    items = [(p.style.name, p.text) for p in doc.paragraphs if p.style.name.startswith("List Bullet")]
    assert items == [("List Bullet", "1. Intro"), ("List Bullet 2", "1.1 Basics")]


def test_docx_course_follows_chapter_sequential_vertical_order(tmp_path):
    headings = _headings(_docx(tmp_path))
    course = headings[headings.index(("Heading 1", "Example Course")) :]
    assert course[:6] == [
        ("Heading 1", "Example Course"),
        ("Heading 2", "1. Intro"),
        ("Heading 3", "1.1 Basics"),
        ("Heading 4", "Lecture 1"),
        ("Heading 5", "Lecture 1"),
        ("Heading 4", "Exercise 1"),
    ]


def test_docx_chapter_summary_and_lecture_text_are_written(tmp_path):
    text = "\n".join(p.text for p in _docx(tmp_path).paragraphs)
    assert "Subject intro." in text
    assert "Text." in text
    assert "alert(1)" not in text


def test_docx_renders_reading_page_inline_under_its_title(tmp_path):
    headings = _headings(_docx(tmp_path))
    reading = [h for h in headings if h == ("Heading 5", "Reading 1")]
    assert len(reading) == 1
    assert ("Heading 4", "Reading 1") in headings


def test_docx_renders_a_page_linked_twice_every_time(tmp_path):
    site = _site()
    second = "https://ocw.tudelft.nl/courses/example/subjects/2-more/"
    site.add_home([(SUBJECT_URL, "1. Intro"), (second, "2. More")])
    site.add_subject_page(
        second,
        f"""
        <ul class="activities"><li>
        <h4><span class="course-title">2.1 Again</span></h4>
        <ul><li><a class="icon icon--lecture" href="{LECTURE_URL}">Lecture 1</a></li></ul>
        </li></ul>
        """,
    )
    doc = _docx(tmp_path, site)
    assert [p.text for p in doc.paragraphs].count("Text.") == 2


def test_docx_embeds_course_images_and_links_downloads_into_assets(tmp_path):
    site = _site()
    site.add_static_file("diagram.png", _png_bytes())
    doc = _docx(tmp_path, site)
    assert len(doc.inline_shapes) == 1
    targets = [
        doc.part.rels[link.get(qn("r:id"))].target_ref
        for p in doc.paragraphs
        for link in p._p.findall(qn("w:hyperlink"))
    ]
    assert "assets/Chapter.pdf" in targets


def test_docx_has_no_relative_page_links_or_iframes(tmp_path):
    doc = _docx(tmp_path)
    targets = [rel.target_ref for rel in doc.part.rels.values() if rel.is_external]
    assert not any(t.endswith(".html") or t.startswith("..") for t in targets)
    assert "https://www.youtube.com/embed/abc123" in targets


def test_docx_empty_vertical_gets_the_placeholder(tmp_path):
    doc = _docx(tmp_path, _site(empty_lecture=True))
    assert any("No content was found on the source page." in p.text for p in doc.paragraphs)


def test_video_component_is_only_a_link_in_docx_mode():
    builder = ZipBuilder(_site().course())
    comp = {"type": "video", "display_name": "Intro", "src": "https://v.example/x"}
    assert "<iframe" not in builder._render_docx(comp, PurePosixPath("a.html"))
    assert "<iframe" in builder._render(comp, PurePosixPath("a.html"))


def test_self_referencing_page_does_not_loop():
    builder = ZipBuilder(_site().course(keep_all_content=True))
    builder.course.parse()
    builder._plan()
    page = builder._by_url[_key(READING_URL)]
    link = {"type": "reading_link", "reading_url": READING_URL}
    page.components = [link]
    html = builder._render_docx(link, PurePosixPath("v.html"))
    assert html.count("(see above)") == 1
    assert html.count("<h5") == 1
