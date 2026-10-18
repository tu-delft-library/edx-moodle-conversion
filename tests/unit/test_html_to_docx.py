import logging

import pytest
from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from ocw.converter.html_to_docx import HtmlToDocx


@pytest.fixture
def assets(tmp_path):
    folder = tmp_path / "assets"
    folder.mkdir()
    return folder


@pytest.fixture
def scratch(tmp_path):
    folder = tmp_path / "scratch"
    folder.mkdir()
    return folder


@pytest.fixture
def doc():
    return Document()


@pytest.fixture
def walker(doc, assets, scratch):
    return HtmlToDocx(doc, assets, scratch)


def _links(doc):
    found = []
    for paragraph in doc.paragraphs:
        for link in paragraph._p.findall(qn("w:hyperlink")):
            target = doc.part.rels[link.get(qn("r:id"))].target_ref
            found.append((target, "".join(t.text for t in link.iter(qn("w:t")))))
    return found


def _image(path, size=(40, 30), fmt=None):
    Image.new("RGB", size, "red").save(path, format=fmt)


def test_headings_and_paragraphs_keep_style_and_text(walker, doc):
    walker.add("<h2>Title</h2><p>Hello <strong>bold</strong> and <em>it</em></p>")
    assert [(p.style.name, p.text) for p in doc.paragraphs] == [
        ("Heading 2", "Title"),
        ("Normal", "Hello bold and it"),
    ]
    runs = {r.text: (r.bold, r.italic) for r in doc.paragraphs[1].runs}
    assert runs["bold"] == (True, None)
    assert runs["it"] == (None, True)


def test_heading_offset_shifts_levels_and_caps_at_nine(walker, doc):
    walker.add("<h1>A</h1><h6>B</h6>", heading_offset=3)
    assert [p.style.name for p in doc.paragraphs] == ["Heading 4", "Heading 9"]


def test_structural_headings_ignore_the_offset(walker, doc):
    walker.add("<h2 data-structural>Chapter</h2><h1>Inside</h1>", heading_offset=3)
    assert [p.style.name for p in doc.paragraphs] == ["Heading 2", "Heading 4"]


def test_whitespace_collapses_and_no_empty_paragraphs(walker, doc):
    walker.add("<div>\n  <p>  one \n two  </p>\n\n  <p>   </p>\n</div>")
    assert [p.text for p in doc.paragraphs] == ["one two"]


def test_loose_text_between_blocks_becomes_its_own_paragraph(walker, doc):
    walker.add("<div>intro<p>body</p>outro</div>")
    assert [p.text for p in doc.paragraphs] == ["intro", "body", "outro"]


def test_br_adds_a_line_break(walker, doc):
    walker.add("<p>a<br>b</p>")
    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0]._p.xpath(".//w:br")


def test_nested_lists_use_list_styles_per_level(walker, doc):
    walker.add("<ul><li>one<ul><li>deep<ol><li>deeper</li></ol></li></ul></li><li>two</li></ul>")
    assert [(p.style.name, p.text) for p in doc.paragraphs] == [
        ("List Bullet", "one"),
        ("List Bullet 2", "deep"),
        ("List Number 3", "deeper"),
        ("List Bullet", "two"),
    ]


def test_paragraph_inside_list_item_stays_in_the_item(walker, doc):
    walker.add("<ul><li><p>first</p><p>second</p></li></ul>")
    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0].text == "first\nsecond"
    assert doc.paragraphs[0]._p.xpath(".//w:br")


def test_external_link_becomes_hyperlink_with_its_text(walker, doc):
    walker.add('<p>see <a href="https://example.org/x">the docs</a>.</p>')
    assert _links(doc) == [("https://example.org/x", "the docs")]
    assert doc.paragraphs[0].text == "see the docs."


def test_link_without_text_shows_its_url(walker, doc):
    walker.add('<p><a href="https://example.org/x"></a></p>')
    assert _links(doc) == [("https://example.org/x", "https://example.org/x")]


@pytest.mark.parametrize("href", ["../../assets/Chapter.pdf", "../assets/Chapter.pdf", "assets/Chapter.pdf"])
def test_local_asset_links_are_relative_to_the_zip_root(walker, doc, href):
    walker.add(f'<p><a href="{href}">chapter</a></p>')
    assert _links(doc) == [("assets/Chapter.pdf", "chapter")]


def test_page_fragment_is_kept_on_local_links(walker, doc):
    walker.add('<p><a href="../assets/Chapter.pdf#page=3">p3</a></p>')
    assert _links(doc) == [("assets/Chapter.pdf#page=3", "p3")]


def test_fragment_only_link_is_plain_text(walker, doc):
    walker.add('<p><a href="#top">back to top</a></p>')
    assert _links(doc) == []
    assert doc.paragraphs[0].text == "back to top"


def test_link_around_image_embeds_the_image(walker, doc, assets):
    _image(assets / "pic.png")
    walker.add('<p><a href="https://example.org"><img src="../assets/pic.png"></a></p>')
    assert len(doc.inline_shapes) == 1
    assert _links(doc) == []


def test_iframe_is_written_as_a_link_to_its_source(walker, doc):
    walker.add('<div><iframe src="https://www.youtube.com/embed/abc"></iframe></div>')
    assert _links(doc) == [("https://www.youtube.com/embed/abc", "https://www.youtube.com/embed/abc")]


def test_table_cells_rows_and_colspan(walker, doc):
    walker.add(
        '<table><tr><th colspan="2">Head</th></tr>'
        "<tr><td>a</td><td><p>b1</p><p>b2</p></td></tr></table>"
    )
    table = doc.tables[0]
    assert len(table.rows) == 2 and len(table.columns) == 2
    assert table.cell(0, 0).text == "Head"
    assert table.cell(1, 0).text == "a"
    assert table.cell(1, 1).text == "b1\nb2"


def test_table_cell_has_no_leading_empty_paragraph(walker, doc):
    walker.add("<table><tr><td><p>x</p></td></tr></table>")
    assert len(doc.tables[0].cell(0, 0).paragraphs) == 1


def test_table_without_rows_is_skipped(walker, doc):
    walker.add("<table></table><p>after</p>")
    assert doc.tables == []
    assert [p.text for p in doc.paragraphs] == ["after"]


def test_details_summary_is_bold_and_body_follows(walker, doc):
    walker.add("<details><summary>Question</summary><p>Answer</p></details>")
    assert [p.text for p in doc.paragraphs] == ["Question", "Answer"]
    assert doc.paragraphs[0].runs[0].bold is True
    assert doc.paragraphs[1].runs[0].bold is None


def test_pre_keeps_lines_in_monospace(walker, doc):
    walker.add("<pre>a = 1\nb = 2</pre>")
    assert len(doc.paragraphs) == 1
    assert {r.font.name for r in doc.paragraphs[0].runs if r.text.strip()} == {"Courier New"}
    assert doc.paragraphs[0]._p.xpath(".//w:br")


def test_sup_and_sub_are_set(walker, doc):
    walker.add("<p>x<sup>up</sup> H<sub>down</sub>O</p>")
    runs = {r.text: r.font for r in doc.paragraphs[0].runs}
    assert runs["up"].superscript is True
    assert runs["down"].subscript is True
    assert not runs["x"].superscript and not runs["x"].subscript


def test_scripts_styles_and_comments_are_dropped(walker, doc):
    walker.add("<p>keep<script>alert(1)</script><style>p{}</style><!-- note --></p>")
    assert [p.text for p in doc.paragraphs] == ["keep"]


def test_illegal_xml_characters_are_removed(walker, doc):
    walker.add("<p>a\x0bb</p>")
    assert doc.paragraphs[0].text == "ab"


def test_unknown_tag_keeps_its_text_and_is_logged_once(walker, doc, caplog):
    with caplog.at_level(logging.INFO, logger="ocw.html_to_docx"):
        walker.add("<p><blink>one</blink> <blink>two</blink></p>")
    assert doc.paragraphs[0].text == "one two"
    assert caplog.text.count("Unhandled <blink>") == 1


def test_known_inline_tags_are_not_logged(walker, caplog):
    with caplog.at_level(logging.INFO, logger="ocw.html_to_docx"):
        walker.add("<p><span>a</span><u>b</u></p>")
    assert "Unhandled" not in caplog.text


def test_png_is_embedded(walker, doc, assets):
    _image(assets / "diagram.png")
    walker.add('<p><img src="../../assets/diagram.png" alt="a diagram"></p>')
    assert len(doc.inline_shapes) == 1
    assert doc.paragraphs[0].text == ""


def test_small_raster_is_still_embedded(walker, doc, assets):
    _image(assets / "icon.png", size=(8, 8))
    walker.add('<p><img src="../assets/icon.png"></p>')
    assert len(doc.inline_shapes) == 1


def test_wide_image_is_capped_to_the_page_width(walker, doc, assets):
    _image(assets / "wide.png", size=(3000, 100))
    walker.add('<p><img src="../assets/wide.png"></p>')
    shape = doc.inline_shapes[0]
    assert shape.width == walker._usable_width
    assert shape.height == pytest.approx(shape.width * 100 / 3000, abs=2)


def test_webp_is_converted_to_png_in_scratch_not_assets(walker, doc, assets, scratch):
    _image(assets / "photo.webp", fmt="WEBP")
    walker.add('<p><img src="../assets/photo.webp"></p>')
    assert len(doc.inline_shapes) == 1
    assert (scratch / "photo.webp.png").is_file()
    assert sorted(p.name for p in assets.iterdir()) == ["photo.webp"]


def test_svg_is_ignored_without_marker_or_text(walker, doc, assets, caplog):
    (assets / "icon.svg").write_text("<svg xmlns='http://www.w3.org/2000/svg'/>")
    with caplog.at_level(logging.DEBUG, logger="ocw.html_to_docx"):
        walker.add('<p>text <img src="../assets/icon.svg" alt="icon"></p>')
    assert len(doc.inline_shapes) == 0
    assert _links(doc) == []
    assert doc.paragraphs[0].text == "text"
    assert (assets / "icon.svg").is_file()


def test_unreadable_raster_gets_a_filename_marker_and_warning(walker, doc, assets, caplog):
    (assets / "bad.png").write_bytes(b"not an image")
    with caplog.at_level(logging.WARNING, logger="ocw.html_to_docx"):
        walker.add('<p><img src="../assets/bad.png" alt="secret alt"></p>')
    assert _links(doc) == [("assets/bad.png", "figure: bad.png")]
    assert "secret alt" not in doc.paragraphs[0].text
    assert "bad.png" in caplog.text


def test_remote_image_gets_a_marker_to_its_url_and_warning(walker, doc, caplog):
    with caplog.at_level(logging.WARNING, logger="ocw.html_to_docx"):
        walker.add('<p><img src="https://example.org/pics/remote.png"></p>')
    assert _links(doc) == [("https://example.org/pics/remote.png", "figure: remote.png")]
    assert "not a local asset" in caplog.text


def test_missing_local_file_gets_a_marker(walker, doc):
    walker.add('<p><img src="../assets/gone.png"></p>')
    assert _links(doc) == [("assets/gone.png", "figure: gone.png")]


def test_alt_text_is_never_written(walker, doc, assets):
    _image(assets / "pic.png")
    walker.add('<p><img src="../assets/pic.png" alt="DECORATIVE ALT"></p>')
    assert "DECORATIVE" not in "".join(p.text for p in doc.paragraphs)


def test_page_break_heading_starts_a_new_page(walker, doc):
    walker.add("<h2 data-page-break>Chapter</h2><h2>Other</h2>")
    assert doc.paragraphs[0].paragraph_format.page_break_before is True
    assert doc.paragraphs[1].paragraph_format.page_break_before is None


_WIDGET = (
    '<div class="wpb_raw_code"><div class="TUD_box_download" data-downloadid="x">'
    '<details class="tud-button video-download"><summary><img src="../../assets/icon_video.png">'
    '<span>Video</span></summary><ul class="TUD">'
    '<li><a href="https://delftxdownloads.tudelft.nl/x-video.mp4">High<span>1080p</span></a></li></ul></details>'
    '<details class="tud-button subtitle-download"><ul class="TUD">'
    '<li><a href="https://delftxdownloads.tudelft.nl/x-video.srt">SRT English</a></li></ul></details>'
    '<details class="tud-button other-download"><ul class="TUD">'
    '<li><a href="../../assets/x-slides.pdf">Slides<span>PDF</span></a></li></ul></details>'
    '<p class="TUD_subtext no-download">Sorry but there don\'t seem to be any downloads..</p>'
    "</div></div>"
)


def test_download_widget_becomes_a_type_file_link_table(walker, doc):
    walker.add(_WIDGET)
    rows = [[c.text for c in row.cells] for row in doc.tables[0].rows]
    assert rows == [
        ["Type", "File", "Link"],
        ["Video", "High 1080p", "https://delftxdownloads.tudelft.nl/x-video.mp4"],
        ["Subtitles", "SRT English", "https://delftxdownloads.tudelft.nl/x-video.srt"],
        ["Other", "Slides PDF", "assets/x-slides.pdf"],
    ]
    assert len(doc.tables[0].cell(1, 2).paragraphs) == 1
    assert doc.paragraphs == [] and len(doc.inline_shapes) == 0


def test_download_widget_without_real_links_is_dropped_with_a_warning(walker, doc, caplog):
    with caplog.at_level(logging.WARNING, logger="ocw.html_to_docx"):
        walker.add(
            '<div class="TUD_box_download"><div class="tud-button video-download"><ul>'
            '<li><a href="#">High</a></li></ul></div><p>Sorry but there are no downloads</p></div><p>after</p>'
        )
    assert "no real links" in caplog.text
    assert doc.tables == []
    assert [p.text for p in doc.paragraphs] == ["after"]
